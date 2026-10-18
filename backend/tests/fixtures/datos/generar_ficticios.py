"""Genera los archivos ficticios de esta carpeta. Es determinista: siempre produce lo mismo.

Uso (desde la raíz del repositorio, con las dependencias de backend/requirements-dev.txt):

    python datos-de-prueba/generar_ficticios.py

Todo lo que aparece aquí es inventado: no hay personas ni identificadores reales.
"""

import csv
import datetime as dt
import random
import zipfile
from pathlib import Path

import openpyxl
import xlwt
from openpyxl.styles import Alignment, Font

CARPETA = Path(__file__).resolve().parent
SEMILLA = 20241018
FECHA_FIJA = dt.datetime(2024, 10, 18, 12, 0, 0)


def guardar_xlsx(libro: openpyxl.Workbook, ruta: Path) -> None:
    """Guarda el libro sin fechas del momento, para que dos ejecuciones den el mismo archivo."""
    libro.properties.created = libro.properties.modified = FECHA_FIJA
    libro.save(ruta)
    with zipfile.ZipFile(ruta) as origen:
        entradas = [(entrada.filename, origen.read(entrada)) for entrada in origen.infolist()]
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as destino:
        for nombre, datos in entradas:
            destino.writestr(zipfile.ZipInfo(nombre, FECHA_FIJA.timetuple()[:6]), datos, zipfile.ZIP_DEFLATED)


# --- pacientes-ficticios.xlsx ---------------------------------------------------------------

DIAGNOSTICOS = [
    # (nombre, peso, probabilidad de mujer, probabilidad de cirugía abierta)
    ("Colecistitis aguda", 30, 0.65, 0.15),
    ("Colecistitis crónica litiásica", 25, 0.75, 0.05),
    ("Apendicitis aguda", 20, 0.45, 0.20),
    ("Hernia inguinal", 12, 0.15, 0.40),
    ("Hernia umbilical", 6, 0.60, 0.35),
    ("Pancreatitis biliar", 5, 0.55, 0.30),
    ("Eventración", 2, 0.65, 0.85),
]
COMORBILIDADES = [
    ("Ninguna", 35),
    ("HTA", 20),
    ("Diabetes", 12),
    ("HTA y Diabetes", 13),
    ("Obesidad", 8),
    ("HTA y Obesidad", 5),
    ("Asma", 3),
    ("Cardiopatía", 4),
]
COMPLICACIONES = [
    "Infección de herida",
    "Íleo prolongado",
    "Sangrado leve",
    "Fiebre postoperatoria",
    "Seroma",
    "Dolor persistente",
]
ENCABEZADOS_PACIENTES = [
    "No.",
    "Sexo",
    "Edad",
    "Diagnóstico",
    "Tipo de cirugía",
    "Conversión",
    "Estancia hospitalaria (días)",
    "Comorbilidades",
    "Tiempo quirúrgico (min)",
    "Complicaciones",
    None,  # columna final sin nombre y vacía, como en el archivo real
]


def generar_pacientes(azar: random.Random, total: int = 200) -> openpyxl.Workbook:
    sin_edad = set(azar.sample(range(total), 3))
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Hoja1"
    hoja.append(ENCABEZADOS_PACIENTES)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for numero in range(1, total + 1):
        diagnostico, _, p_mujer, p_abierta = azar.choices(DIAGNOSTICOS, weights=[d[1] for d in DIAGNOSTICOS])[0]
        sexo = "Femenino" if azar.random() < p_mujer else "Masculino"
        edad = None if numero - 1 in sin_edad else min(92, max(18, round(azar.gauss(52, 17))))
        abierta = azar.random() < p_abierta
        convertida = (not abierta) and azar.random() < 0.10
        if abierta:
            conversion = "No aplica"
        else:
            conversion = "Sí" if convertida else "No"
        estancia = max(1, round(azar.gauss(6 if abierta or convertida else 3, 1.8)))
        tiempo = max(30, round(azar.gauss(150 if abierta or convertida else 85, 25)))
        comorbilidad = azar.choices([c[0] for c in COMORBILIDADES], weights=[c[1] for c in COMORBILIDADES])[0]
        complicacion = azar.choice(COMPLICACIONES) if azar.random() < 0.2 else None
        hoja.append(
            [
                numero,
                sexo,
                edad,
                diagnostico,
                "Abierta" if abierta else "Laparoscópica",
                conversion,
                estancia,
                comorbilidad,
                tiempo,
                complicacion,
            ]
        )
    # Celdas con formato pero sin valor: la columna "existe" en la hoja aunque esté vacía
    for fila in range(1, total + 2):
        hoja.cell(row=fila, column=len(ENCABEZADOS_PACIENTES)).alignment = Alignment(horizontal="center")
    hoja.freeze_panes = "A2"
    for letra, ancho in zip("ABCDEFGHIJ", (6, 12, 8, 30, 16, 12, 28, 18, 24, 24), strict=True):
        hoja.column_dimensions[letra].width = ancho
    libro.create_sheet("Hoja2")
    libro.create_sheet("Hoja3")
    return libro


# --- ficticio-punto-y-coma.csv ---------------------------------------------------------------


def generar_csv(azar: random.Random, ruta: Path, total: int = 60) -> None:
    municipios = ["Cañada Verde", "Piedra Añeja", "Valle Alto", "Monte Sombrío", "Río Manso"]
    diagnosticos = ["Hipertensión arterial", "Diabetes tipo 2", "Asma bronquial", "Gastritis crónica", "Anemia ferropénica"]
    with ruta.open("w", encoding="cp1252", newline="") as archivo:
        escritor = csv.writer(archivo, delimiter=";", lineterminator="\n")
        escritor.writerow(["Código", "Sexo", "Edad", "Peso (kg)", "Talla (m)", "Municipio de residencia", "Diagnóstico", "Fecha de ingreso"])
        for numero in range(1, total + 1):
            peso = round(azar.gauss(70, 12), 1)
            talla = round(azar.gauss(1.65, 0.09), 2)
            ingreso = dt.date(2024, 1, 1) + dt.timedelta(days=azar.randrange(0, 180))
            escritor.writerow(
                [
                    f"C-{numero:03d}",
                    azar.choice(["Femenino", "Masculino"]),
                    azar.randrange(18, 90),
                    str(peso).replace(".", ","),
                    str(talla).replace(".", ","),
                    azar.choice(municipios),
                    azar.choice(diagnosticos),
                    ingreso.isoformat(),
                ]
            )


# --- ficticio-antiguo.xls --------------------------------------------------------------------


def generar_xls(azar: random.Random, ruta: Path, total: int = 50) -> None:
    libro = xlwt.Workbook(encoding="utf-8")
    hoja = libro.add_sheet("Consultas")
    estilo_fecha = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for columna, encabezado in enumerate(["Código", "Servicio", "Edad", "Fecha de ingreso", "Presión sistólica", "Fumador"]):
        hoja.write(0, columna, encabezado)
    for fila in range(1, total + 1):
        hoja.write(fila, 0, f"A-{fila:03d}")
        hoja.write(fila, 1, azar.choice(["Cirugía general", "Medicina interna", "Pediatría", "Ginecología"]))
        hoja.write(fila, 2, azar.randrange(1, 95))
        hoja.write(fila, 3, dt.date(2023, 1, 1) + dt.timedelta(days=azar.randrange(0, 300)), estilo_fecha)
        hoja.write(fila, 4, azar.randrange(90, 180))
        hoja.write(fila, 5, azar.choice(["Sí", "No"]))
    libro.save(str(ruta))


# --- ficticio-con-suciedad.xlsx ----------------------------------------------------------------


def generar_suciedad(azar: random.Random, ruta: Path, total: int = 40) -> None:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(["Paciente", "Sexo", "Edad", "Peso (kg)", "Servicio"])
    variantes_sexo = ["Masculino", "masculino", "Masculino ", "Femenino", "Femenino", "femenino"]
    filas_vacias = {8, 15, 16, 30}
    numero = 0
    fila_actual = 1
    while numero < total:
        fila_actual += 1
        if fila_actual in filas_vacias:
            hoja.append([None] * 5)
            continue
        numero += 1
        edad: object = azar.randrange(20, 85)
        if numero in (5, 22):
            edad = "N/D"
        elif numero % 6 == 0:
            edad = str(edad)  # número guardado como texto
        peso = round(azar.gauss(72, 11), 1)
        hoja.append(
            [
                f"P-{numero:03d}",
                azar.choice(variantes_sexo),
                edad,
                str(peso).replace(".", ",") if numero % 5 == 0 else peso,
                azar.choice(["Urgencias", "Consulta externa", "Hospitalización"]),
            ]
        )
    guardar_xlsx(libro, ruta)


# --- ficticio-titulo-en-primera-fila.xlsx -------------------------------------------------------


def generar_titulo(azar: random.Random, ruta: Path, total: int = 30) -> None:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Enero"
    hoja.append(["Registro ficticio de consultas - enero"])
    hoja.append(["Folio", "Motivo de consulta", "Edad", "Tiempo de espera (min)"])
    for numero in range(1, total + 1):
        hoja.append(
            [
                f"F-{numero:03d}",
                azar.choice(["Control", "Dolor abdominal", "Revisión de herida", "Fiebre"]),
                azar.randrange(5, 90),
                azar.randrange(5, 120),
            ]
        )
    guardar_xlsx(libro, ruta)


def main() -> None:
    guardar_xlsx(generar_pacientes(random.Random(SEMILLA)), CARPETA / "pacientes-ficticios.xlsx")
    generar_csv(random.Random(SEMILLA + 1), CARPETA / "ficticio-punto-y-coma.csv")
    generar_xls(random.Random(SEMILLA + 2), CARPETA / "ficticio-antiguo.xls")
    generar_suciedad(random.Random(SEMILLA + 3), CARPETA / "ficticio-con-suciedad.xlsx")
    generar_titulo(random.Random(SEMILLA + 4), CARPETA / "ficticio-titulo-en-primera-fila.xlsx")


if __name__ == "__main__":
    main()
