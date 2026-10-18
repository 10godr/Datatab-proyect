"""Funciones de apoyo para las pruebas: crean archivos ficticios en memoria."""

import csv
import io
import re
import zipfile
from collections.abc import Callable, Sequence
from typing import Any

import openpyxl
import xlwt

from app.config import Configuracion
from app.services.columnas import PerfilColumna
from app.services.file_reader import TablaLeida, leer_archivo


def crear_xlsx(filas: Sequence[Sequence[Any]], hojas_extra: dict[str, Sequence[Sequence[Any]]] | None = None) -> bytes:
    libro = openpyxl.Workbook()
    libro.active.title = "Hoja1"
    for fila in filas:
        libro.active.append(list(fila))
    for nombre, contenido in (hojas_extra or {}).items():
        hoja = libro.create_sheet(nombre)
        for fila in contenido:
            hoja.append(list(fila))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def crear_xls(filas: Sequence[Sequence[Any]]) -> bytes:
    libro = xlwt.Workbook(encoding="utf-8")
    hoja = libro.add_sheet("Hoja1")
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            if valor is not None:
                hoja.write(i, j, valor)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def crear_csv(
    filas: Sequence[Sequence[Any]], separador: str = ",", codificacion: str = "utf-8", fin_linea: str = "\n"
) -> bytes:
    texto = io.StringIO()
    csv.writer(texto, delimiter=separador, lineterminator=fin_linea).writerows(filas)
    return texto.getvalue().encode(codificacion)


def leer(nombre: str, contenido: bytes, config: Configuracion | None = None) -> TablaLeida:
    return leer_archivo(nombre, contenido, config or Configuracion())


def leer_filas(filas: Sequence[Sequence[Any]], config: Configuracion | None = None) -> TablaLeida:
    """Pasa las filas por el lector completo (como un .xlsx) y devuelve la tabla leída."""
    return leer("datos.xlsx", crear_xlsx(filas), config)


def columna(leida: TablaLeida, nombre: str) -> tuple[Any, PerfilColumna]:
    perfil = next(p for p in leida.perfiles if p.nombre == nombre)
    return leida.tabla[nombre], perfil


def codigos_avisos(leida: TablaLeida) -> list[str]:
    return [a["codigo"] for a in leida.avisos]


def reescribir_zip(contenido: bytes, cambios: Callable[[str, bytes], bytes | None]) -> bytes:
    """Copia un zip aplicando `cambios(nombre, datos)` a cada entrada (None = la quita)."""
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(contenido)) as origen, zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
        for entrada in origen.infolist():
            datos = cambios(entrada.filename, origen.read(entrada))
            if datos is not None:
                destino.writestr(entrada.filename, datos)
    return salida.getvalue()


def a_cadenas_compartidas(contenido: bytes) -> bytes:
    """Convierte los textos en línea de un .xlsx de openpyxl a textos compartidos, como los guarda Excel."""
    textos: list[str] = []

    def compartir(celda: re.Match) -> str:
        if celda[3] not in textos:
            textos.append(celda[3])
        return f'<c r="{celda[1]}"{celda[2]} t="s"><v>{textos.index(celda[3])}</v></c>'

    def cambios(nombre: str, datos: bytes) -> bytes:
        if nombre == "xl/worksheets/sheet1.xml":
            patron = r'<c r="([A-Z]+[0-9]+)"((?: s="[0-9]+")?) t="inlineStr"><is><t[^>]*>(.*?)</t></is></c>'
            return re.sub(patron, compartir, datos.decode()).encode()
        if nombre == "xl/_rels/workbook.xml.rels":
            relacion = (
                '<Relationship Id="rIdSS" Target="sharedStrings.xml" Type="http://schemas.openxmlformats.org/'
                'officeDocument/2006/relationships/sharedStrings"/></Relationships>'
            )
            return datos.decode().replace("</Relationships>", relacion).encode()
        if nombre == "[Content_Types].xml":
            tipo = (
                '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
                'officedocument.spreadsheetml.sharedStrings+xml"/></Types>'
            )
            return datos.decode().replace("</Types>", tipo).encode()
        return datos

    convertido = reescribir_zip(contenido, cambios)
    tabla = "".join(f'<si><t xml:space="preserve">{t}</t></si>' for t in textos)
    return agregar_al_zip(
        convertido,
        "xl/sharedStrings.xml",
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{tabla}</sst>'.encode(),
    )


def agregar_al_zip(contenido: bytes, nombre: str, datos: bytes) -> bytes:
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(contenido)) as origen, zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
        for entrada in origen.infolist():
            destino.writestr(entrada.filename, origen.read(entrada))
        destino.writestr(nombre, datos)
    return salida.getvalue()
