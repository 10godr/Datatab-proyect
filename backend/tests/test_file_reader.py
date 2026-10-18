import dataclasses
import datetime as dt
import io
import re
import time
import tracemalloc
import zipfile

import openpyxl
import pytest
import xlwt
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app.config import Configuracion
from app.errors import ErrorDataTab
from app.services.file_reader import FIRMA_OLE, limpiar_nombre
from ayudas import (
    a_cadenas_compartidas,
    agregar_al_zip,
    codigos_avisos,
    columna,
    crear_csv,
    crear_xls,
    crear_xlsx,
    leer,
    leer_filas,
    reescribir_zip,
)
from conftest import DATOS_DE_PRUEBA

SECRETO = "SECRETO-CELDA-8841"


def error_de(nombre: str, contenido: bytes, config: Configuracion | None = None) -> ErrorDataTab:
    with pytest.raises(ErrorDataTab) as capturado:
        leer(nombre, contenido, config)
    return capturado.value


def aviso(leida, codigo: str) -> dict:
    return next(a for a in leida.avisos if a["codigo"] == codigo)


# --- Nombre y extensión ---------------------------------------------------------------------


@pytest.mark.parametrize(
    "entrada, esperado",
    [
        ("pacientes.xlsx", "pacientes.xlsx"),
        ("../../etc/passwd.xlsx", "passwd.xlsx"),
        ("..\\..\\Windows\\datos.csv", "datos.csv"),
        ("C:\\Users\\yo\\Mis datos.xls", "Mis datos.xls"),
        ("/var/tmp/a.xlsx", "a.xlsx"),
        ("con\x00nulo\x07y\x1bcontrol\r\n.xlsx", "connuloycontrol.xlsx"),
        ("bidi\u202egnp.xlsx", "bidignp.xlsx"),
        ("  espacios.csv  ", "espacios.csv"),
        ("", "archivo"),
        (None, "archivo"),
        ("../", "archivo"),
        ("línea\u2028nueva.csv", "líneanueva.csv"),
    ],
)
def test_limpiar_nombre(entrada, esperado):
    assert limpiar_nombre(entrada) == esperado


def test_nombre_largo_se_recorta_conservando_la_extension():
    nombre = limpiar_nombre("a" * 500 + ".xlsx")
    assert len(nombre) == 120 and nombre.endswith(".xlsx")
    assert len(limpiar_nombre("b" * 500)) == 120


@pytest.mark.parametrize("nombre", ["datos.pdf", "datos", "datos.xlsx.exe", "datos.docx", "datos.", "", "xlsx", "datos.xlsm"])
def test_formato_no_soportado(nombre):
    error = error_de(nombre, b"contenido")
    assert error.codigo == "FORMATO_NO_SOPORTADO" and error.estado == 415
    assert error.detalle == {"formatos": ["xlsx", "xls", "csv"]}


def test_la_extension_no_distingue_mayusculas():
    assert leer("DATOS.XLSX", crear_xlsx([["a"], [1]])).formato == "xlsx"
    assert leer("Datos.Csv", crear_csv([["a"], [1]])).formato == "csv"


# --- xlsx ---------------------------------------------------------------------------------------------


def test_xlsx_basico():
    leida = leer_filas([["Sexo", "Edad"], ["Femenino", 34], ["Masculino", 41], ["Femenino", None]])
    assert leida.formato == "xlsx" and leida.nombre == "datos.xlsx"
    assert list(leida.tabla.columns) == ["Sexo", "Edad"] and len(leida.tabla) == 3
    assert [(p.nombre, p.tipo, p.vacios) for p in leida.perfiles] == [("Sexo", "texto", 0), ("Edad", "numerica", 1)]
    assert leida.avisos == []


def test_las_columnas_son_texto_con_valores_normalizados_y_nan_como_vacio():
    leida = leer_filas([["Edad", "Nota"], [30.0, "  a   b "], [None, "x"], [31.5, None]])
    assert leida.tabla["Edad"].tolist()[0] == "30" and leida.tabla["Edad"].tolist()[2] == "31.5"
    assert leida.tabla["Nota"].tolist()[0] == "a b"
    assert leida.tabla["Edad"].isna().tolist() == [False, True, False]
    assert str(leida.tabla["Edad"].dtype) == "str"


def test_tipos_de_excel_bool_fecha_y_hora():
    filas = [
        ["Bool", "Fecha", "FechaHora", "Numero"],
        [True, dt.date(2024, 5, 1), dt.datetime(2024, 5, 1, 13, 45), 3],
        [False, dt.datetime(2024, 5, 2), dt.datetime(2024, 5, 2, 8, 0), 4],
    ]
    leida = leer_filas(filas)
    assert leida.tabla["Bool"].tolist() == ["Verdadero", "Falso"]
    assert leida.tabla["Fecha"].tolist() == ["2024-05-01", "2024-05-02"]
    assert leida.tabla["FechaHora"].tolist() == ["2024-05-01 13:45", "2024-05-02 08:00"]
    assert [p.tipo for p in leida.perfiles] == ["texto", "fecha", "fecha", "numerica"]


def test_solo_se_lee_la_primera_hoja_y_se_avisa_si_las_otras_tienen_datos():
    contenido = crear_xlsx([["a"], [1]], hojas_extra={"Otra": [["x"], [2]], "Vacia": []})
    leida = leer("datos.xlsx", contenido)
    assert list(leida.tabla.columns) == ["a"]
    mensaje = aviso(leida, "VARIAS_HOJAS")
    assert mensaje["nivel"] == "info" and "2 hojas con datos" in mensaje["mensaje"] and "«Hoja1»" in mensaje["mensaje"]


def test_las_hojas_vacias_no_avisan():
    leida = leer("datos.xlsx", crear_xlsx([["a"], [1]], hojas_extra={"Hoja2": [], "Hoja3": [[None, None]]}))
    assert "VARIAS_HOJAS" not in codigos_avisos(leida)


def test_se_descartan_columnas_vacias_pero_no_las_que_tienen_nombre():
    leida = leer_filas([["A", None, "C", "D"], [1, None, None, 4], [2, None, None, 5]])
    assert list(leida.tabla.columns) == ["A", "C", "D"]  # la B no tiene nada: se descarta
    assert (columna(leida, "C")[1].tipo, columna(leida, "C")[1].vacios, columna(leida, "C")[1].distintos) == ("texto", 2, 0)


def test_se_descartan_filas_vacias_y_se_cuentan_las_intercaladas():
    leida = leer_filas([["A", "B"], [1, 2], [None, None], ["", "  "], [3, 4], [None, None], [None, None]])
    assert len(leida.tabla) == 2
    aviso_filas = aviso(leida, "FILAS_VACIAS_OMITIDAS")
    assert aviso_filas["nivel"] == "info" and "2 filas vacías" in aviso_filas["mensaje"]  # las del final no cuentan


def test_una_sola_fila_vacia_en_singular():
    leida = leer_filas([["A"], [1], [None], [2]])
    assert aviso(leida, "FILAS_VACIAS_OMITIDAS")["mensaje"] == "Se omitió 1 fila vacía."


def test_sin_filas_vacias_no_hay_aviso():
    assert "FILAS_VACIAS_OMITIDAS" not in codigos_avisos(leer_filas([["A"], [1], [2]]))


def test_celdas_con_solo_espacios_son_vacias():
    leida = leer_filas([["A", "B"], ["x", "   "], ["y", "\t\n"]])
    assert columna(leida, "B")[1].vacios == 2


def test_columna_sin_nombre_se_llama_columna_n_segun_su_posicion():
    leida = leer_filas([["A", None, "C", None, "E"], [1, 2, 3, 4, 5]])
    assert list(leida.tabla.columns) == ["A", "Columna 2", "C", "Columna 4", "E"]


def test_la_columna_final_sin_nombre_y_vacia_no_cuenta_para_los_encabezados_dudosos():
    leida = leer_filas([["A", "B", None, None, None], [1, 2, None, None, None]])
    assert list(leida.tabla.columns) == ["A", "B"] and leida.avisos == []


def test_encabezados_dudosos_si_mas_de_la_mitad_no_tiene_nombre():
    # 3 de 4 sin nombre
    leida = leer_filas([["Título", None, None, None], ["a", "b", "c", "d"]])
    dudoso = aviso(leida, "ENCABEZADOS_DUDOSOS")
    assert dudoso["nivel"] == "atencion" and dudoso["mensaje"] == "La primera fila no parece tener nombres de columna."
    # la mitad exacta no es "más de la mitad"
    assert "ENCABEZADOS_DUDOSOS" not in codigos_avisos(leer_filas([["A", "B", None, None], [1, 2, 3, 4]]))


def test_nombres_de_columna_se_limpian_y_los_numeros_son_texto():
    leida = leer_filas([["  Edad   en años ", 2020, dt.date(2024, 1, 1)], [1, 2, 3]])
    assert list(leida.tabla.columns) == ["Edad en años", "2020", "2024-01-01"]


def test_nombres_repetidos_se_renombran_y_se_avisa():
    leida = leer_filas([["Edad", "Edad", "Peso", "Edad"], [1, 2, 3, 4]])
    assert list(leida.tabla.columns) == ["Edad", "Edad (2)", "Peso", "Edad (3)"]
    repetidas = aviso(leida, "COLUMNAS_REPETIDAS")
    assert "«Edad»" in repetidas["mensaje"] and "«Edad (2)»" in repetidas["mensaje"]


def test_renombrar_no_choca_con_un_nombre_que_ya_existe():
    leida = leer_filas([["Edad", "Edad (2)", "Edad"], [1, 2, 3]])
    assert list(leida.tabla.columns) == ["Edad", "Edad (2)", "Edad (3)"]
    assert len(set(leida.tabla.columns)) == 3


def test_nombres_que_difieren_solo_en_mayusculas_no_son_repetidos():
    leida = leer_filas([["Edad", "edad"], [1, 2]])
    assert list(leida.tabla.columns) == ["Edad", "edad"] and "COLUMNAS_REPETIDAS" not in codigos_avisos(leida)


def test_numeros_guardados_como_texto_cuentan_como_los_numeros():
    leida = leer_filas([["Edad"], [45], ["45"], [" 45 "], ["45.0"], [30], [31], [32], [33], [34], [35]])
    serie, perfil = columna(leida, "Edad")
    assert perfil.tipo == "numerica" and perfil.distintos == 7
    assert serie.value_counts()["45"] == 4


def test_valores_parecidos_se_avisan_en_la_lectura():
    leida = leer_filas([["Sexo"], ["Masculino"], ["masculino"], ["Masculino "], ["Femenino"]])
    assert "«Masculino» y «masculino»" in aviso(leida, "VALORES_PARECIDOS")["mensaje"]
    serie, perfil = columna(leida, "Sexo")
    assert perfil.distintos == 3 and serie.value_counts()["Masculino"] == 2  # "Masculino " y "Masculino" son uno


def test_archivo_con_solo_encabezados_o_sin_nada_es_vacio():
    for filas in ([["A", "B"]], [[None, None], [None, None]], []):
        error = error_de("datos.xlsx", crear_xlsx(filas))
        assert error.codigo == "ARCHIVO_VACIO" and error.estado == 422


def test_contenido_de_cero_bytes_es_vacio():
    for nombre in ("a.xlsx", "a.xls", "a.csv"):
        assert error_de(nombre, b"").codigo == "ARCHIVO_VACIO"


# --- xlsx: casos de celdas, hojas y formatos que escribe Excel -----------------------------------------


def hoja_xml(contenido: bytes, parte: str = "xl/worksheets/sheet1.xml") -> str:
    with zipfile.ZipFile(io.BytesIO(contenido)) as comprimido:
        return comprimido.read(parte).decode()


def con_hoja_cambiada(contenido: bytes, cambiar) -> bytes:
    """Aplica `cambiar` (str -> str) al XML de la primera hoja."""
    return reescribir_zip(
        contenido, lambda nombre, datos: cambiar(datos.decode()).encode() if nombre == "xl/worksheets/sheet1.xml" else datos
    )


def guardar(libro: openpyxl.Workbook) -> bytes:
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_las_celdas_de_error_cuentan_como_vacias():
    contenido = crear_xlsx([["A", "B", "C"], ["x", 5, "z"], ["y", 6, "w"]])
    con_error = con_hoja_cambiada(contenido, lambda xml: xml.replace('<c r="B2" t="n"><v>5</v></c>', '<c r="B2" t="e"><v>#DIV/0!</v></c>'))
    assert "#DIV/0!" in hoja_xml(con_error)
    leida = leer("datos.xlsx", con_error)
    assert leida.tabla["B"].isna().tolist() == [True, False]
    assert "#DIV/0!" not in leida.tabla["B"].dropna().tolist()


def test_un_texto_que_parece_un_error_de_excel_sigue_siendo_texto():
    contenido = crear_xlsx([["Dato"], ["#N/A"], ["#N/A"], ["ok"]])  # openpyxl guarda "#N/A" como celda de error
    como_texto = con_hoja_cambiada(contenido, lambda xml: xml.replace('t="e"><v>#N/A</v>', 't="inlineStr"><is><t>#N/A</t></is>'))
    assert "t=\"e\"" not in hoja_xml(como_texto) and hoja_xml(contenido).count('t="e"') == 2
    assert leer("datos.xlsx", como_texto).tabla["Dato"].tolist() == ["#N/A", "#N/A", "ok"]
    assert leer("datos.xlsx", contenido).tabla["Dato"].tolist() == ["ok"]  # como error de Excel, en cambio, es vacío


def test_fecha_con_numero_de_serie_fuera_de_rango_cuenta_como_vacia():
    libro = openpyxl.Workbook()
    libro.active.append(["Dia", "Otro"])
    libro.active.append([dt.date(2024, 5, 1), 1])
    libro.active.append([10**12, 2])
    libro.active["A3"].number_format = "yyyy-mm-dd"
    with pytest.warns(UserWarning, match="marked as a date"):  # openpyxl avisa; DataTab lo trata como vacío
        leida = leer("datos.xlsx", guardar(libro))
    assert leida.tabla["Dia"].tolist()[0] == "2024-05-01" and leida.tabla["Dia"].isna().tolist() == [False, True]


def test_fechas_horas_y_duraciones_de_excel():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Fecha", "FechaHora", "Hora", "Duracion"])
    hoja.append([dt.date(2024, 5, 1), dt.datetime(2024, 5, 1, 13, 45, 30), dt.time(9, 5), dt.timedelta(hours=30, minutes=15)])
    hoja["D2"].number_format = "[h]:mm:ss"
    leida = leer("datos.xlsx", guardar(libro) + b"")
    fila = {nombre: leida.tabla[nombre].tolist()[0] for nombre in leida.tabla.columns}
    assert fila == {"Fecha": "2024-05-01", "FechaHora": "2024-05-01 13:45", "Hora": "09:05", "Duracion": "30:15"}


def test_calendario_1904():
    libro = openpyxl.Workbook()
    libro.epoch = CALENDAR_MAC_1904
    libro.active.append(["Dia"])
    libro.active.append([dt.datetime(2024, 5, 1)])
    assert leer("datos.xlsx", guardar(libro)).tabla["Dia"].tolist() == ["2024-05-01"]


def test_celdas_y_filas_sueltas_lejos_de_las_demas():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja["A1"], hoja["C1"], hoja["AA1"] = "a", "c", "lejos"
    hoja["B4"] = 4
    hoja["A6"] = "fin"
    leida = leer("datos.xlsx", guardar(libro))
    assert list(leida.tabla.columns) == ["a", "Columna 2", "c", "lejos"]
    assert len(leida.tabla) == 2
    # Las filas 2, 3 y 5 no existen en el archivo, pero en la hoja son filas vacías entre filas con datos
    assert aviso(leida, "FILAS_VACIAS_OMITIDAS")["mensaje"] == "Se omitieron 3 filas vacías."


def test_dimension_declarada_que_no_es_cierta_no_recorta_la_hoja():
    contenido = crear_xlsx([["A", "B"]] + [[i, i * 2] for i in range(1, 21)])
    recortada = con_hoja_cambiada(contenido, lambda xml: re.sub(r'<dimension ref="[^"]*"\s*/>', '<dimension ref="A1:A2"/>', xml))
    assert 'ref="A1:A2"' in hoja_xml(recortada)
    leida = leer("datos.xlsx", recortada)
    assert list(leida.tabla.columns) == ["A", "B"] and len(leida.tabla) == 20


def test_hoja_sin_dimension_declarada():
    contenido = crear_xlsx([["A", "B"], [1, 2], [3, 4]])
    sin_dimension = con_hoja_cambiada(contenido, lambda xml: re.sub(r'<dimension ref="[^"]*"\s*/>', "", xml))
    assert "<dimension" not in hoja_xml(sin_dimension)
    assert len(leer("datos.xlsx", sin_dimension).tabla) == 2


def test_formulas_dan_el_valor_que_guardo_excel_y_vacio_si_no_lo_guardo():
    contenido = crear_xlsx([["A", "B", "C", "D"], [1, 2, "=A2+B2", '="a"&"b"']])
    sin_valor = leer("datos.xlsx", contenido)  # openpyxl no calcula: no hay valor guardado
    assert sin_valor.tabla["C"].isna().tolist() == [True] and sin_valor.tabla["D"].isna().tolist() == [True]
    fila_con_valores = (  # como la guarda Excel: la fórmula y el último resultado calculado
        '<row r="2"><c r="A2"><v>1</v></c><c r="B2"><v>2</v></c>'
        '<c r="C2"><f>A2+B2</f><v>3</v></c><c r="D2" t="str"><f>"a"&amp;"b"</f><v>ab</v></c></row>'
    )
    con_valor = con_hoja_cambiada(contenido, lambda xml: re.sub(r'<row r="2">.*?</row>', fila_con_valores, xml))
    assert "<v>ab</v>" in hoja_xml(con_valor)
    leida = leer("datos.xlsx", con_valor)
    assert (leida.tabla["C"].tolist(), leida.tabla["D"].tolist()) == (["3"], ["ab"])


def test_textos_compartidos_como_los_guarda_excel():
    original = crear_xlsx([["Sexo", "Nota"], ["Femenino", "a & b <c>"], ["Masculino", "  espacios  "], ["Femenino", "Ñ"]])
    compartido = a_cadenas_compartidas(original)
    assert "xl/sharedStrings.xml" in zipfile.ZipFile(io.BytesIO(compartido)).namelist()
    assert "inlineStr" not in hoja_xml(compartido)
    leida = leer("datos.xlsx", compartido)
    assert leida.tabla["Sexo"].tolist() == ["Femenino", "Masculino", "Femenino"]
    assert leida.tabla["Nota"].tolist() == ["a & b <c>", "espacios", "Ñ"]


def test_textos_con_formato_y_con_guia_fonetica():
    libro = openpyxl.Workbook()
    libro.active.append(["Titulo", "Otro"])
    libro.active["A2"] = CellRichText("Hola ", TextBlock(InlineFont(b=True), "mundo"), " fin")
    libro.active["B2"] = "normal"
    assert leer("datos.xlsx", guardar(libro)).tabla["Titulo"].tolist() == ["Hola mundo fin"]

    xml = (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<si><t>Dato</t></si><si><t>Dos</t></si>'
        '<si><r><t>Ho</t></r><r><rPr><b/></rPr><t>la</t></r><rPh sb="0" eb="1"><t>GUIA</t></rPh></si>'
        '<si><t>sin formato</t><phoneticPr fontId="1"/></si></sst>'
    )
    base = a_cadenas_compartidas(crear_xlsx([["Dato", "Dos"], ["x", "y"]]))
    sustituido = reescribir_zip(base, lambda nombre, datos: xml.encode() if nombre == "xl/sharedStrings.xml" else datos)
    cambiado = con_hoja_cambiada(sustituido, lambda hoja: hoja.replace("<v>2</v>", "<v>X</v>").replace("<v>3</v>", "<v>Y</v>"))
    leida = leer("datos.xlsx", reescribir_zip(cambiado, lambda n, d: d.replace(b"<v>X</v>", b"<v>2</v>").replace(b"<v>Y</v>", b"<v>3</v>")))
    assert leida.tabla["Dato"].tolist() == ["Hola"] and leida.tabla["Dos"].tolist() == ["sin formato"]


def test_las_hojas_de_grafico_no_son_hojas_de_datos():
    libro = openpyxl.Workbook()
    datos = libro.active
    datos.title = "Datos"
    datos.append(["a"])
    datos.append([1])
    grafico = BarChart()
    grafico.add_data(Reference(datos, min_col=1, min_row=1, max_row=2), titles_from_data=True)
    libro.create_chartsheet("Grafico", 0).add_chart(grafico)  # la primera pestaña es un gráfico
    leida = leer("datos.xlsx", guardar(libro))
    assert list(leida.tabla.columns) == ["a"] and "VARIAS_HOJAS" not in codigos_avisos(leida)


def test_xlsx_de_la_hoja_en_modo_escritura_sin_dimension():
    libro = openpyxl.Workbook(write_only=True)
    hoja = libro.create_sheet("Datos")
    for fila in (["x", "y"], [1, 2], [3, 4]):
        hoja.append(fila)
    leida = leer("datos.xlsx", guardar(libro))
    assert list(leida.tabla.columns) == ["x", "y"] and len(leida.tabla) == 2


# --- xls ------------------------------------------------------------------------------------------------


def test_xls_basico_con_fechas_y_bool():
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Datos")
    estilo = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM")
    for j, titulo in enumerate(["Nombre", "Edad", "Ingreso", "Activo"]):
        hoja.write(0, j, titulo)
    hoja.write(1, 0, "Ana")
    hoja.write(1, 1, 34)
    hoja.write(1, 2, dt.datetime(2024, 5, 1, 13, 45), estilo)
    hoja.write(1, 3, True)
    hoja.write(2, 0, "Luis")
    hoja.write(2, 1, 41.5)
    hoja.write(2, 2, dt.datetime(2024, 5, 2), estilo)
    hoja.write(2, 3, False)
    libro.add_sheet("Otra").write(0, 0, "dato")
    salida = io.BytesIO()
    libro.save(salida)
    leida = leer("viejo.xls", salida.getvalue())
    assert leida.formato == "xls"
    assert leida.tabla["Nombre"].tolist() == ["Ana", "Luis"]
    assert leida.tabla["Edad"].tolist() == ["34", "41.5"]
    assert leida.tabla["Ingreso"].tolist() == ["2024-05-01 13:45", "2024-05-02"]
    assert leida.tabla["Activo"].tolist() == ["Verdadero", "Falso"]
    assert [p.tipo for p in leida.perfiles] == ["texto", "numerica", "fecha", "texto"]
    assert "VARIAS_HOJAS" in codigos_avisos(leida)


def test_xls_filas_vacias_y_columnas_sin_nombre():
    leida = leer("viejo.xls", crear_xls([["A", None, "C"], [1, None, 3], [None, None, None], [4, None, 6]]))
    assert list(leida.tabla.columns) == ["A", "C"] and len(leida.tabla) == 2
    assert "Se omitió 1 fila vacía." == aviso(leida, "FILAS_VACIAS_OMITIDAS")["mensaje"]


def test_xls_vacio_es_vacio():
    assert error_de("viejo.xls", crear_xls([["A", "B"]])).codigo == "ARCHIVO_VACIO"


# --- csv ----------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("separador", [",", ";", "\t"])
def test_csv_detecta_el_separador(separador):
    leida = leer("datos.csv", crear_csv([["Sexo", "Edad", "Peso"], ["F", 34, "72.5"], ["M", 41, "80"]], separador))
    assert list(leida.tabla.columns) == ["Sexo", "Edad", "Peso"] and len(leida.tabla) == 2


def test_csv_punto_y_coma_con_decimales_con_coma():
    contenido = "Peso (kg);Talla (m)\n72,5;1,68\n80;1,75\n".encode("cp1252")
    leida = leer("datos.csv", contenido)
    assert list(leida.tabla.columns) == ["Peso (kg)", "Talla (m)"]
    assert leida.tabla["Peso (kg)"].tolist() == ["72.5", "80"]
    assert all(p.tipo == "numerica" for p in leida.perfiles) and leida.perfiles[1].minimo == 1.68


def test_csv_con_comas_decimales_entre_comillas():
    leida = leer("datos.csv", b'Peso,Nota\n"72,5","a, b"\n"80","c"\n')
    assert leida.tabla["Peso"].tolist() == ["72.5", "80"] and leida.tabla["Nota"].tolist() == ["a, b", "c"]


@pytest.mark.parametrize("codificacion", ["utf-8", "utf-8-sig", "cp1252"])
def test_csv_codificaciones_con_tildes_y_enie(codificacion):
    contenido = crear_csv([["Diagnóstico", "Año"], ["Niño con tos", "Peñón"]], ";", codificacion)
    leida = leer("datos.csv", contenido)
    assert list(leida.tabla.columns) == ["Diagnóstico", "Año"]
    assert leida.tabla["Diagnóstico"].tolist() == ["Niño con tos"]  # sin BOM pegado al primer nombre


def test_csv_con_fin_de_linea_windows_y_lineas_vacias():
    leida = leer("datos.csv", b"A,B\r\n1,2\r\n\r\n3,4\r\n\r\n")
    assert len(leida.tabla) == 2 and aviso(leida, "FILAS_VACIAS_OMITIDAS")["mensaje"] == "Se omitió 1 fila vacía."


def test_csv_con_salto_de_linea_dentro_de_comillas():
    leida = leer("datos.csv", b'A,B\n"linea 1\nlinea 2",x\ny,z\n')
    assert leida.tabla["A"].tolist() == ["linea 1 linea 2", "y"]


def test_csv_con_filas_de_distinto_largo():
    leida = leer("datos.csv", b"a,b\n1,2,3\n4\n5,6\n")
    assert list(leida.tabla.columns) == ["a", "b", "Columna 3"]
    assert leida.tabla["Columna 3"].tolist()[0] == "3" and leida.tabla["b"].isna().tolist() == [False, True, False]


def test_csv_con_separador_final_en_cada_linea():
    leida = leer("datos.csv", b"a;b;\n1;2;\n3;4;\n")
    assert list(leida.tabla.columns) == ["a", "b"]


def test_csv_de_una_sola_columna():
    leida = leer("datos.csv", b"Edad\n30\n31\n")
    assert list(leida.tabla.columns) == ["Edad"] and len(leida.tabla) == 2


def test_csv_solo_encabezados_es_vacio():
    assert error_de("datos.csv", b"a,b,c\n").codigo == "ARCHIVO_VACIO"
    assert error_de("datos.csv", b"\n\n\n").codigo == "ARCHIVO_VACIO"


def test_csv_todo_es_texto_y_el_tipo_lo_decide_datatab():
    leida = leer("datos.csv", b"Id,Fecha,Edad\n001,2024-05-01,30\n002,2024-05-02,31\n")
    assert [p.tipo for p in leida.perfiles] == ["numerica", "fecha", "numerica"]
    assert leida.tabla["Id"].tolist() == ["1", "2"]  # en una columna numérica "001" es el número 1


# --- Contenido que no corresponde a la extensión -------------------------------------------------------------


def test_xlsx_que_en_realidad_es_texto():
    error = error_de("datos.xlsx", b"esto no es un excel, solo texto")
    assert (error.codigo, error.estado) == ("CONTENIDO_NO_COINCIDE", 415)


def test_xlsx_que_es_un_zip_de_otra_cosa():
    contenido = reescribir_zip(crear_xlsx([["a"], [1]]), lambda nombre, datos: None if nombre.startswith("xl/") else datos)
    assert error_de("datos.xlsx", contenido).codigo == "CONTENIDO_NO_COINCIDE"
    docx = io.BytesIO()
    with zipfile.ZipFile(docx, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "<w:document/>")
    assert error_de("datos.xlsx", docx.getvalue()).codigo == "CONTENIDO_NO_COINCIDE"


def test_xlsx_cifrado_simulado_con_la_firma_ole():
    error = error_de("datos.xlsx", FIRMA_OLE + bytes(4096))
    assert error.codigo == "CONTENIDO_NO_COINCIDE" and error.estado == 415
    assert "contraseña" in error.mensaje and "contraseña" in error.sugerencia and ".xlsx" in error.sugerencia


def test_xls_que_es_html():
    html = f"<html><body><table><tr><td>{SECRETO}</td></tr></table></body></html>".encode()
    for contenido in (html, b"\xef\xbb\xbf  \n" + html, b"<?xml version='1.0'?><Workbook/>"):
        error = error_de("informe.xls", contenido)
        assert error.codigo == "CONTENIDO_NO_COINCIDE" and "página web" in error.mensaje
        assert SECRETO not in error.mensaje + (error.sugerencia or "")


def test_xls_que_es_otra_cosa():
    assert error_de("datos.xls", crear_xlsx([["a"], [1]])).codigo == "CONTENIDO_NO_COINCIDE"
    assert error_de("datos.xls", b"texto cualquiera").codigo == "CONTENIDO_NO_COINCIDE"


def test_csv_binario():
    binario = bytes(range(256)) * 20
    for contenido in (binario, b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR", crear_xlsx([["a"], [1]]), "a,b\n1,2".encode("utf-16")):
        error = error_de("datos.csv", contenido)
        assert error.codigo == "CONTENIDO_NO_COINCIDE" and error.estado == 415


def test_csv_con_bytes_que_no_se_pueden_decodificar():
    # 0x81 no existe en Windows-1252 y no es UTF-8 válido
    assert error_de("datos.csv", b"a,b\n\x81,2\n").codigo == "CONTENIDO_NO_COINCIDE"


# --- Archivos dañados -----------------------------------------------------------------------------------------------


def test_xlsx_corrupto():
    valido = crear_xlsx([["a", "b"], [1, 2]] * 50)
    for contenido in (valido[: len(valido) // 2], b"PK\x03\x04" + b"basura" * 100, b"PK"):
        error = error_de("datos.xlsx", contenido)
        assert (error.codigo, error.estado) == ("ARCHIVO_ILEGIBLE", 422), contenido[:10]


def test_xlsx_con_xml_roto_no_filtra_su_contenido():
    xml_roto = f'<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>{SECRETO}</t></is></c></row><</sheetData>'

    def cambios(nombre, datos):
        return xml_roto.encode() if nombre == "xl/worksheets/sheet1.xml" else datos

    error = error_de("datos.xlsx", reescribir_zip(crear_xlsx([["a"], [1]]), cambios))
    assert error.codigo == "ARCHIVO_ILEGIBLE"
    assert SECRETO not in repr(error.__dict__)


def test_xlsx_con_indice_de_texto_fuera_de_rango():
    def cambios(nombre, datos):
        if nombre == "xl/worksheets/sheet1.xml":
            return datos.decode().replace('t="inlineStr"><is><t>a</t></is>', 't="s"><v>99</v>').encode()
        return datos

    assert error_de("datos.xlsx", reescribir_zip(crear_xlsx([["a"], [1]]), cambios)).codigo == "ARCHIVO_ILEGIBLE"


def test_xls_corrupto():
    error = error_de("datos.xls", FIRMA_OLE + bytes(2000))
    assert error.codigo == "ARCHIVO_ILEGIBLE"
    valido = crear_xls([["a", "b"], [1, 2]] * 30)
    assert error_de("datos.xls", valido[: len(valido) // 2]).codigo == "ARCHIVO_ILEGIBLE"


def test_csv_con_un_campo_gigante_es_ilegible():
    assert error_de("datos.csv", b"a,b\n" + b"x" * 200_000 + b",2\n").codigo == "ARCHIVO_ILEGIBLE"


# --- xlsx con DTD: ataques de expansión de entidades ----------------------------------------------------------------------

HOJA_MINIMA = (
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
    '<row r="1"><c r="A1" t="inlineStr"><is><t>&lol9;</t></is></c></row></sheetData></worksheet>'
)


def dtd_de_entidades_anidadas(niveles: int = 9) -> str:
    """«Mil millones de risas»: cada entidad repite diez veces la anterior (10^9 repeticiones de «lol»)."""
    definiciones = '<!ENTITY lol0 "lol">'
    for nivel in range(1, niveles + 1):
        definiciones += f'<!ENTITY lol{nivel} "' + f"&lol{nivel - 1};" * 10 + '">'
    return f"<!DOCTYPE lolz [{definiciones}]>"


def con_dtd(parte_xml: bytes, dtd: str, antes_del_dtd: str = "") -> bytes:
    """Inserta un DTD justo después de la declaración `<?xml ...?>` de una parte (o al principio)."""
    texto = parte_xml.decode()
    declaracion = texto[: texto.index("?>") + 2] if texto.startswith("<?xml") else ""
    return (declaracion + antes_del_dtd + dtd + texto[len(declaracion) :]).encode()


def xlsx_con_la_hoja(xml: bytes) -> bytes:
    return reescribir_zip(crear_xlsx([["a"], [1]]), lambda nombre, datos: xml if nombre == "xl/worksheets/sheet1.xml" else datos)


def no_debe_abrirse_el_libro(monkeypatch):
    """La defensa va antes de abrir el libro: si se llega a `load_workbook`, la prueba falla."""
    from app.services import file_reader

    def prohibido(*_, **__):
        pytest.fail("se llegó a abrir un libro con DTD")

    monkeypatch.setattr(file_reader, "load_workbook", prohibido)


def test_una_bomba_de_entidades_se_rechaza_al_instante_y_sin_consumir_memoria(monkeypatch):
    bomba = xlsx_con_la_hoja(('<?xml version="1.0"?>' + dtd_de_entidades_anidadas() + HOJA_MINIMA).encode())
    no_debe_abrirse_el_libro(monkeypatch)
    tracemalloc.start()
    try:
        inicio = time.perf_counter()
        error = error_de("datos.xlsx", bomba)
        duracion = time.perf_counter() - inicio
        pico = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert (error.codigo, error.estado) == ("ARCHIVO_ILEGIBLE", 422)
    assert duracion < 0.5 and pico < 2 * 1024 * 1024, (duracion, pico)  # sin la defensa: ~1 s y decenas de MB
    assert "lol" not in repr(error.__dict__)


@pytest.mark.parametrize(
    "dtd, antes_del_dtd, codificacion",
    [
        ("<!DOCTYPE worksheet>", "", "utf-8"),
        ('<!DOCTYPE worksheet [<!ENTITY a "a">]>', "", "utf-8"),
        ('<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>', "", "utf-8"),
        ("<!DOCTYPE worksheet>", "<!-- " + "x" * 20_000 + " -->", "utf-8"),  # tras mucho relleno: pasa de cualquier ventana fija
        ("<!DOCTYPE worksheet>", " " * 20_000, "utf-8"),
        ("<!DOCTYPE worksheet>", "", "utf-16"),  # con otra codificación el texto «<!DOCTYPE» no aparece en los bytes
    ],
    ids=["doctype", "entidad", "entidad-externa", "tras-comentario-largo", "tras-espacios", "utf-16"],
)
def test_cualquier_dtd_se_rechaza_aunque_se_oculte(monkeypatch, dtd, antes_del_dtd, codificacion):
    xml = '<?xml version="1.0"?>' + antes_del_dtd + dtd + HOJA_MINIMA.replace("&lol9;", "a")
    if codificacion == "utf-16":
        xml = '<?xml version="1.0" encoding="utf-16"?>' + antes_del_dtd + dtd + HOJA_MINIMA.replace("&lol9;", "a")
    no_debe_abrirse_el_libro(monkeypatch)
    error = error_de("datos.xlsx", xlsx_con_la_hoja(xml.encode(codificacion)))
    assert (error.codigo, error.estado) == ("ARCHIVO_ILEGIBLE", 422)


@pytest.mark.parametrize(
    "parte",
    [
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/worksheets/sheet1.xml",
    ],
)
def test_el_dtd_se_rechaza_en_cualquier_parte_del_archivo(monkeypatch, parte):
    contenido = reescribir_zip(
        crear_xlsx([["a"], [1]]),
        lambda nombre, datos: con_dtd(datos, dtd_de_entidades_anidadas()) if nombre == parte else datos,
    )
    no_debe_abrirse_el_libro(monkeypatch)
    assert error_de("datos.xlsx", contenido).codigo == "ARCHIVO_ILEGIBLE"


def test_el_dtd_se_rechaza_aunque_la_parte_no_se_llame_xml(monkeypatch):
    contenido = agregar_al_zip(crear_xlsx([["a"], [1]]), "xl/extra.bin", con_dtd(HOJA_MINIMA.encode(), dtd_de_entidades_anidadas()))
    no_debe_abrirse_el_libro(monkeypatch)
    assert error_de("datos.xlsx", contenido).codigo == "ARCHIVO_ILEGIBLE"


def test_las_partes_que_no_son_xml_no_se_confunden_con_un_dtd():
    # Un archivo adjunto con bytes cualquiera (aquí, un PNG que menciona «<!DOCTYPE» en un metadato) no es un DTD
    imagen = b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 8 + b"<!DOCTYPE <!ENTITY"
    contenido = agregar_al_zip(crear_xlsx([["a"], [1]]), "xl/media/imagen1.png", imagen)
    assert leer("datos.xlsx", contenido).tabla["a"].tolist() == ["1"]


# --- Límites ----------------------------------------------------------------------------------------------------------


def test_zip_bomb_se_rechaza_antes_de_abrirlo():
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/worksheets/sheet1.xml", bytes(201 * 1024 * 1024))
    bomba = salida.getvalue()
    assert len(bomba) < 1024 * 1024  # pesa menos de 1 MB y descomprime a más de 200 MB
    error = error_de("datos.xlsx", bomba)
    assert (error.codigo, error.estado) == ("TABLA_DEMASIADO_GRANDE", 422)
    assert error.detalle == {"max_descomprimido_mb": 200} and "descomprimido" in error.mensaje


def test_limite_de_descomprimido_configurable():
    contenido = crear_xlsx([["a"], [1]])
    pequeno = dataclasses.replace(Configuracion(), max_descomprimido_mb=1)
    assert leer("datos.xlsx", contenido, pequeno).formato == "xlsx"  # el archivo normal cabe de sobra
    relleno = reescribir_zip(contenido, lambda n, d: d + b" " * 2_000_000 if n == "docProps/app.xml" else d)
    assert error_de("datos.xlsx", relleno, pequeno).codigo == "TABLA_DEMASIADO_GRANDE"


def test_limite_de_filas():
    config = dataclasses.replace(Configuracion(), max_filas=3)
    assert len(leer_filas([["A"], [1], [2], [3]], config).tabla) == 3
    error = error_de("datos.xlsx", crear_xlsx([["A"], [1], [2], [3], [4]]), config)
    assert (error.codigo, error.estado) == ("TABLA_DEMASIADO_GRANDE", 422) and error.detalle == {"max_filas": 3}


def test_las_filas_vacias_no_cuentan_para_el_limite_de_filas():
    config = dataclasses.replace(Configuracion(), max_filas=2)
    assert len(leer_filas([["A"], [1], [None], [None], [None], [2]], config).tabla) == 2


def test_limite_de_columnas():
    config = dataclasses.replace(Configuracion(), max_columnas=3)
    assert len(leer_filas([["A", "B", "C"], [1, 2, 3]], config).perfiles) == 3
    error = error_de("datos.xlsx", crear_xlsx([["A", "B", "C", "D"], [1, 2, 3, 4]]), config)
    assert error.codigo == "TABLA_DEMASIADO_GRANDE" and error.detalle == {"max_columnas": 3}


def test_las_columnas_vacias_no_cuentan_para_el_limite_de_columnas():
    config = dataclasses.replace(Configuracion(), max_columnas=3)
    fila = ["A", None, None, None, None, "B"]
    assert list(leer_filas([fila, [1, None, None, None, None, 2]], config).tabla.columns) == ["A", "B"]


def test_limite_de_celdas():
    config = dataclasses.replace(Configuracion(), max_celdas=6)
    assert leer_filas([["A", "B"], [1, 2], [3, 4], [5, 6]], config).tabla.shape == (3, 2)
    error = error_de("datos.xlsx", crear_xlsx([["A", "B"], [1, 2], [3, 4], [5, 6], [7, 8]]), config)
    assert error.codigo == "TABLA_DEMASIADO_GRANDE" and error.detalle == {"max_celdas": 6}


def test_limites_tambien_en_csv_y_xls():
    config = dataclasses.replace(Configuracion(), max_filas=1)
    assert error_de("datos.csv", b"a\n1\n2\n", config).codigo == "TABLA_DEMASIADO_GRANDE"
    assert error_de("datos.xls", crear_xls([["a"], [1], [2]]), config).codigo == "TABLA_DEMASIADO_GRANDE"


# --- Ningún mensaje contiene datos --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "nombre, contenido",
    [
        ("a.xlsx", crear_xlsx([[SECRETO, "b"], [SECRETO, 2], [SECRETO, 3]])),
        ("a.xls", crear_xls([[SECRETO, "b"], [SECRETO, 2], [SECRETO, 3]])),
        ("a.csv", f"{SECRETO},b\n{SECRETO},2\n{SECRETO},3\n".encode()),
        ("a.xlsx", f"{SECRETO}".encode()),
        ("a.xls", f"<html>{SECRETO}</html>".encode()),
        ("a.csv", f"{SECRETO}\x00".encode()),
        ("a.xlsx", b"PK\x03\x04" + SECRETO.encode()),
        ("a.xls", FIRMA_OLE + SECRETO.encode()),
    ],
)
def test_los_errores_nunca_incluyen_contenido_de_celdas(nombre, contenido):
    config = dataclasses.replace(Configuracion(), max_filas=1)
    for configuracion in (Configuracion(), config):
        try:
            leer(nombre, contenido, configuracion)
        except ErrorDataTab as error:
            assert SECRETO not in repr(error.__dict__), error.codigo
            assert all(ruta not in repr(error.__dict__) for ruta in ("/home", "/tmp", "site-packages", "Traceback"))


# --- Los archivos ficticios del repositorio ----------------------------------------------------------------------------------


def leer_ficticio(nombre: str):
    return leer(nombre, (DATOS_DE_PRUEBA / nombre).read_bytes())


def test_pacientes_ficticios():
    leida = leer_ficticio("pacientes-ficticios.xlsx")
    assert len(leida.tabla) == 200 and leida.avisos == []
    assert list(leida.tabla.columns) == [
        "No.", "Sexo", "Edad", "Diagnóstico", "Tipo de cirugía", "Conversión",
        "Estancia hospitalaria (días)", "Comorbilidades", "Tiempo quirúrgico (min)", "Complicaciones",
    ]  # fmt: skip
    tipos = {p.nombre: p.tipo for p in leida.perfiles}
    assert tipos["Edad"] == "numerica" and tipos["Sexo"] == "texto" and tipos["Comorbilidades"] == "texto"
    edad = columna(leida, "Edad")[1]
    assert (edad.vacios, edad.minimo, edad.maximo) == (3, 18, 92)
    assert columna(leida, "Complicaciones")[1].vacios > 100
    assert "HTA y Diabetes" in columna(leida, "Comorbilidades")[0].tolist()


def test_ficticio_con_suciedad():
    leida = leer_ficticio("ficticio-con-suciedad.xlsx")
    assert len(leida.tabla) == 40
    edad = columna(leida, "Edad")[1]
    assert (edad.tipo, edad.no_numericos) == ("numerica", 2)  # "N/D" x2; los números guardados como texto cuentan
    assert columna(leida, "Peso (kg)")[1].tipo == "numerica"
    assert set(codigos_avisos(leida)) == {"FILAS_VACIAS_OMITIDAS", "VALORES_PARECIDOS"}
    assert "4 filas vacías" in aviso(leida, "FILAS_VACIAS_OMITIDAS")["mensaje"]
    assert "«Masculino» y «masculino»" in aviso(leida, "VALORES_PARECIDOS")["mensaje"]


def test_ficticio_titulo_en_primera_fila_dispara_encabezados_dudosos():
    leida = leer_ficticio("ficticio-titulo-en-primera-fila.xlsx")
    assert "ENCABEZADOS_DUDOSOS" in codigos_avisos(leida)
    assert list(leida.tabla.columns)[1:] == ["Columna 2", "Columna 3", "Columna 4"]


def test_ficticio_antiguo_xls():
    leida = leer_ficticio("ficticio-antiguo.xls")
    assert leida.formato == "xls" and len(leida.tabla) == 50 and leida.avisos == []
    assert {p.nombre: p.tipo for p in leida.perfiles}["Fecha de ingreso"] == "fecha"


def test_ficticio_csv_punto_y_coma_windows_1252():
    leida = leer_ficticio("ficticio-punto-y-coma.csv")
    assert leida.formato == "csv" and len(leida.tabla) == 60 and leida.avisos == []
    assert "Municipio de residencia" in leida.tabla.columns and "Código" in leida.tabla.columns
    tipos = {p.nombre: p.tipo for p in leida.perfiles}
    assert tipos["Peso (kg)"] == tipos["Talla (m)"] == "numerica" and tipos["Fecha de ingreso"] == "fecha"
    assert "Cañada Verde" in columna(leida, "Municipio de residencia")[0].tolist()
