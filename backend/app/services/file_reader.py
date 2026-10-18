"""Lectura de archivos Excel y CSV desde bytes en memoria (nunca se escribe en disco)."""

import csv
import datetime as dt
import io
import unicodedata
import zipfile
from collections import Counter
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass, replace
from typing import IO, Any
from xml.parsers import expat

import pandas as pd
import xlrd
from openpyxl import load_workbook

from ..config import Configuracion
from ..errors import ErrorDataTab
from ..formato import miles
from .avisos import Aviso, aviso
from .columnas import PerfilColumna, normalizar_valor, perfilar_columna

FORMATOS = ("xlsx", "xls", "csv")
FIRMA_OLE = bytes.fromhex("D0CF11E0A1B11AE1")  # contenedor de .xls (y de .xlsx con contraseña)
MAX_LONGITUD_NOMBRE = 120
MAX_FILAS_REVISADAS = 10_000  # al buscar datos en las hojas que no se leen
MAX_PROLOGO_XML = 64 * 1024  # hasta aquí se busca un DTD; el prólogo de un XML real mide unos cientos de bytes
MUESTRA_SEPARADOR = 20_000
LINEAS_MUESTRA = 30


@dataclass(frozen=True)
class TablaLeida:
    nombre: str
    formato: str
    tabla: pd.DataFrame  # columnas de texto normalizado; vacío = NaN
    perfiles: list[PerfilColumna]
    avisos: list[Aviso]


@dataclass(frozen=True)
class _Hoja:
    """Las celdas de la primera hoja ya como texto normalizado (None = vacía)."""

    encabezado: list[str | None]
    datos: list[list[str | None]]  # filas con al menos una celda con valor, sin la cola vacía
    columnas_usadas: list[int]  # posiciones (desde 0) con algún valor, encabezado incluido
    vacias_omitidas: int  # filas vacías entre filas con datos
    nombre_hoja: str | None = None
    hojas_con_datos: int = 1


# --- Errores del archivo -----------------------------------------------------------------


def _no_coincide(mensaje: str, sugerencia: str) -> ErrorDataTab:
    return ErrorDataTab("CONTENIDO_NO_COINCIDE", mensaje, sugerencia)


def _con_contrasena() -> ErrorDataTab:
    return _no_coincide(
        "Este archivo parece un Excel protegido con contraseña (o no es del formato que indica su extensión).",
        "Quita la contraseña en Excel, o guárdalo como .xlsx sin protección, y vuelve a subirlo.",
    )


def _ilegible() -> ErrorDataTab:
    return ErrorDataTab(
        "ARCHIVO_ILEGIBLE",
        "No se pudo leer el archivo: parece dañado.",
        "Ábrelo en Excel, guárdalo de nuevo y vuelve a subirlo.",
    )


def _vacio(mensaje: str) -> ErrorDataTab:
    return ErrorDataTab(
        "ARCHIVO_VACIO", mensaje, "Revisa que la primera hoja tenga encabezados y al menos una fila de datos."
    )


def _tabla_demasiado_grande(mensaje: str, detalle: dict[str, int]) -> ErrorDataTab:
    return ErrorDataTab(
        "TABLA_DEMASIADO_GRANDE",
        mensaje,
        "Elimina filas o columnas que no necesites y vuelve a intentarlo.",
        detalle,
    )


# --- Nombre y formato ---------------------------------------------------------------------


def limpiar_nombre(nombre: str | None) -> str:
    """Nombre solo para mostrar: sin rutas, sin caracteres de control y de hasta 120 caracteres."""
    base = (nombre or "").replace("\\", "/").rsplit("/", 1)[-1]
    base = "".join(c for c in base if unicodedata.category(c)[0] != "C" and c not in "  ").strip()
    if len(base) > MAX_LONGITUD_NOMBRE:
        cuerpo, punto, extension = base.rpartition(".")
        if punto and len(extension) <= 10:
            base = cuerpo[: MAX_LONGITUD_NOMBRE - len(extension) - 1] + "." + extension
        else:
            base = base[:MAX_LONGITUD_NOMBRE]
    return base or "archivo"


def _formato_por_extension(nombre: str) -> str:
    extension = nombre.rpartition(".")[2].lower() if "." in nombre else ""
    if extension not in FORMATOS:
        raise ErrorDataTab(
            "FORMATO_NO_SOPORTADO",
            "Este tipo de archivo no es compatible.",
            "Sube un archivo de Excel (.xlsx o .xls) o un CSV (.csv).",
            {"formatos": list(FORMATOS)},
        )
    return extension


# --- Ensamblado de la tabla (común a los tres formatos) ----------------------------------


def _acumular(crudas: Iterable[Sequence[Any]], config: Configuracion) -> _Hoja:
    """Normaliza las filas según llegan, aplica los límites y salta las filas vacías.

    La primera fila son los encabezados. Las filas vacías al final no se cuentan: solo
    importan las que quedan entre filas con datos.
    """
    encabezado: list[str | None] | None = None
    datos: list[list[str | None]] = []
    usadas: set[int] = set()
    pendientes = omitidas = 0
    for cruda in crudas:
        fila = [normalizar_valor(valor) for valor in cruda]
        while fila and fila[-1] is None:
            fila.pop()
        if encabezado is None:
            encabezado = fila
        elif not fila:
            pendientes += 1
            continue
        else:
            omitidas += pendientes
            pendientes = 0
            datos.append(fila)
        usadas.update(i for i, valor in enumerate(fila) if valor is not None)

        if len(datos) > config.max_filas:
            raise _tabla_demasiado_grande(
                f"La tabla tiene más de {miles(config.max_filas)} filas.", {"max_filas": config.max_filas}
            )
        if len(usadas) > config.max_columnas:
            raise _tabla_demasiado_grande(
                f"La tabla tiene más de {miles(config.max_columnas)} columnas con datos.",
                {"max_columnas": config.max_columnas},
            )
        if len(datos) * len(usadas) > config.max_celdas:
            raise _tabla_demasiado_grande(
                f"La tabla tiene más de {miles(config.max_celdas)} celdas.", {"max_celdas": config.max_celdas}
            )

    if encabezado is None or not datos:
        raise _vacio("El archivo no tiene filas con datos.")
    return _Hoja(encabezado, datos, sorted(usadas), omitidas)


def _hay_datos(crudas: Iterable[Sequence[Any]]) -> bool:
    for revisadas, cruda in enumerate(crudas):
        if revisadas >= MAX_FILAS_REVISADAS:
            break
        if any(normalizar_valor(valor) is not None for valor in cruda):
            return True
    return False


def _nombrar_columnas(
    encabezado: list[str | None], posiciones: list[int]
) -> tuple[list[str], int, list[tuple[str, str]]]:
    """Nombres finales, cuántas columnas no tenían nombre y las renombradas (original, nuevo)."""
    nombres: list[str] = []
    usados: set[str] = set()
    ultimo_sufijo: dict[str, int] = {}
    sin_nombre = 0
    renombradas: list[tuple[str, str]] = []
    for posicion in posiciones:
        nombre = encabezado[posicion] if posicion < len(encabezado) else None
        if nombre is None:
            sin_nombre += 1
            nombre = f"Columna {posicion + 1}"
        final = nombre
        if final in usados:
            sufijo = ultimo_sufijo.get(nombre, 1)
            while final in usados:  # "Edad (2)" puede existir ya como nombre propio
                sufijo += 1
                final = f"{nombre} ({sufijo})"
            ultimo_sufijo[nombre] = sufijo
            renombradas.append((nombre, final))
        usados.add(final)
        nombres.append(final)
    return nombres, sin_nombre, renombradas


def _construir(nombre: str, formato: str, hoja: _Hoja) -> TablaLeida:
    nombres, sin_nombre, renombradas = _nombrar_columnas(hoja.encabezado, hoja.columnas_usadas)
    avisos: list[Aviso] = []
    if sin_nombre * 2 > len(nombres):
        avisos.append(aviso("ENCABEZADOS_DUDOSOS", "atencion", "La primera fila no parece tener nombres de columna."))
    if renombradas:
        originales = list(dict.fromkeys(original for original, _ in renombradas))[:3]
        lista = ", ".join(f"«{o}»" for o in originales)
        avisos.append(
            aviso(
                "COLUMNAS_REPETIDAS",
                "atencion",
                f"Hay columnas con el mismo nombre ({lista}). Las repetidas se renombraron, "
                f"por ejemplo «{renombradas[0][1]}».",
            )
        )
    if hoja.hojas_con_datos > 1:
        avisos.append(
            aviso(
                "VARIAS_HOJAS",
                "info",
                f"El archivo tiene {hoja.hojas_con_datos} hojas con datos. "
                f"DataTab solo lee la primera («{hoja.nombre_hoja}»).",
            )
        )
    if hoja.vacias_omitidas:
        cuenta = "Se omitió 1 fila vacía" if hoja.vacias_omitidas == 1 else f"Se omitieron {hoja.vacias_omitidas} filas vacías"
        avisos.append(aviso("FILAS_VACIAS_OMITIDAS", "info", f"{cuenta}."))

    series, perfiles = [], []
    for posicion, nombre_columna in zip(hoja.columnas_usadas, nombres, strict=True):
        valores = [fila[posicion] if posicion < len(fila) else None for fila in hoja.datos]
        serie, perfil, avisos_columna = perfilar_columna(nombre_columna, pd.Series(valores, dtype="str"))
        series.append(serie)
        perfiles.append(perfil)
        avisos.extend(avisos_columna)
    tabla = pd.DataFrame(dict(zip(nombres, series, strict=True)))
    return TablaLeida(nombre, formato, tabla, perfiles, avisos)


# --- xlsx ---------------------------------------------------------------------------------------


class _FinDelPrologo(Exception):
    """Detiene el analizador en cuanto se sabe si el XML declara un DTD."""


def _prologo_xml_seguro(flujo: IO[bytes]) -> bool:
    """True si el XML llega a su elemento raíz sin declarar un DTD (`<!DOCTYPE`, con sus `<!ENTITY`).

    Ningún .xlsx real lleva DTD; es la vía de los ataques de expansión de entidades. Como un DTD solo
    puede ir antes del elemento raíz, basta analizar con expat el principio de la parte, sin expandir
    nada. Así se detecta con cualquier codificación y aunque lo precedan miles de espacios o comentarios.
    Lo que no es XML (imágenes, binarios) no tiene DTD y se da por bueno.
    """
    hay_dtd = False

    def encontrar_dtd(*_: Any) -> None:
        nonlocal hay_dtd
        hay_dtd = True
        raise _FinDelPrologo

    def encontrar_raiz(*_: Any) -> None:
        raise _FinDelPrologo

    analizador = expat.ParserCreate()
    analizador.StartDoctypeDeclHandler = encontrar_dtd
    analizador.StartElementHandler = encontrar_raiz
    datos = flujo.read(MAX_PROLOGO_XML)
    try:
        analizador.Parse(datos, False)
    except _FinDelPrologo:
        return not hay_dtd
    except expat.ExpatError:
        return True
    return len(datos) < MAX_PROLOGO_XML  # un prólogo tan largo no es normal: se rechaza por precaución


def _validar_xlsx(contenido: bytes, config: Configuracion) -> None:
    if contenido.startswith(FIRMA_OLE):
        raise _con_contrasena()
    if not contenido.startswith(b"PK"):
        raise _no_coincide(
            "El contenido del archivo no corresponde a un Excel .xlsx.",
            "Comprueba que sea un archivo de Excel y guárdalo de nuevo como .xlsx.",
        )
    with zipfile.ZipFile(io.BytesIO(contenido)) as comprimido:
        entradas = comprimido.infolist()
        if sum(entrada.file_size for entrada in entradas) > config.max_descomprimido_bytes:
            raise _tabla_demasiado_grande(
                "El archivo ocupa demasiado una vez descomprimido.",
                {"max_descomprimido_mb": config.max_descomprimido_mb},
            )
        nombres = {entrada.filename for entrada in entradas}
        if "[Content_Types].xml" not in nombres or not any(n.startswith("xl/") for n in nombres):
            raise _no_coincide(
                "El contenido del archivo no corresponde a un Excel .xlsx.",
                "Comprueba que sea un archivo de Excel y guárdalo de nuevo como .xlsx.",
            )
        for entrada in entradas:
            with comprimido.open(entrada) as flujo:
                if not _prologo_xml_seguro(flujo):
                    raise _ilegible()


def _filas_xlsx(hoja: Any) -> Iterator[list[Any]]:
    """Valores de las celdas de una hoja, fila a fila. Las celdas de error cuentan como vacías."""
    hoja.reset_dimensions()  # las dimensiones que declara el archivo no siempre son ciertas
    for fila in hoja.iter_rows():
        yield [None if celda.data_type == "e" else celda.value for celda in fila]


def _leer_xlsx(contenido: bytes, config: Configuracion) -> _Hoja:
    _validar_xlsx(contenido, config)
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True, keep_links=False)
    try:
        hojas = libro.worksheets  # las hojas de gráfico no tienen datos y no están en la lista
        if not hojas:
            raise _vacio("El archivo no tiene hojas.")
        primera = _acumular(_filas_xlsx(hojas[0]), config)
        otras = sum(1 for hoja in hojas[1:] if _hay_datos(_filas_xlsx(hoja)))
        return replace(primera, nombre_hoja=hojas[0].title, hojas_con_datos=1 + otras)
    finally:
        libro.close()


# --- xls --------------------------------------------------------------------------------------


def _valor_xls(tipo: int, valor: Any, modo_fecha: int) -> Any:
    if tipo == xlrd.XL_CELL_DATE:
        try:
            *fecha, hora, minuto, segundo = xlrd.xldate_as_tuple(valor, modo_fecha)
        except xlrd.XLDateError:
            return valor
        if fecha == [0, 0, 0]:
            return dt.time(hora, minuto, segundo)
        return dt.datetime(*fecha, hora, minuto, segundo)
    if tipo == xlrd.XL_CELL_BOOLEAN:
        return bool(valor)
    if tipo in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    return valor


def _filas_xls(hoja: xlrd.sheet.Sheet, modo_fecha: int) -> Iterator[list[Any]]:
    for fila in range(hoja.nrows):
        yield [_valor_xls(t, v, modo_fecha) for t, v in zip(hoja.row_types(fila), hoja.row_values(fila), strict=True)]


def _leer_xls(contenido: bytes, config: Configuracion) -> _Hoja:
    if not contenido.startswith(FIRMA_OLE):
        if contenido.lstrip(b"\xef\xbb\xbf \t\r\n").startswith(b"<"):
            raise _no_coincide(
                "Este archivo no es un Excel .xls real: es una página web (HTML o XML) guardada con extensión .xls.",
                "Ábrelo en Excel y guárdalo como .xlsx, o expórtalo como CSV.",
            )
        raise _no_coincide(
            "El contenido del archivo no corresponde a un Excel .xls.",
            "Comprueba que sea un archivo de Excel y guárdalo de nuevo como .xlsx.",
        )
    try:
        libro = xlrd.open_workbook(file_contents=contenido, on_demand=True)
    except xlrd.XLRDError as error:
        if "encrypted" in str(error).lower():
            raise _con_contrasena() from None
        raise
    try:
        if libro.nsheets == 0:
            raise _vacio("El archivo no tiene hojas.")
        hojas = [libro.sheet_by_index(i) for i in range(libro.nsheets)]
        primera = _acumular(_filas_xls(hojas[0], libro.datemode), config)
        otras = sum(1 for hoja in hojas[1:] if _hay_datos(_filas_xls(hoja, libro.datemode)))
        return replace(primera, nombre_hoja=hojas[0].name, hojas_con_datos=1 + otras)
    finally:
        libro.release_resources()


# --- csv --------------------------------------------------------------------------------------


def _decodificar_csv(contenido: bytes) -> str:
    no_es_csv = _no_coincide(
        "Este archivo no parece un CSV de texto.",
        "Si es una hoja de cálculo, súbela como .xlsx o .xls; si es un CSV, guárdalo con codificación UTF-8.",
    )
    if b"\x00" in contenido:  # binario, o texto en UTF-16
        raise no_es_csv
    try:
        return contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return contenido.decode("cp1252")
    except UnicodeDecodeError:
        raise no_es_csv from None


def _detectar_separador(texto: str) -> str:
    """Elige entre `,` `;` y tabulador el que da más líneas con el mismo número de campos."""
    lineas = [linea for linea in texto[:MUESTRA_SEPARADOR].split("\n") if linea.strip()]
    if len(texto) > MUESTRA_SEPARADOR:
        lineas = lineas[:-1]  # la última puede estar cortada
    lineas = lineas[:LINEAS_MUESTRA]
    mejor, mejor_puntuacion = ",", (0, 0)
    for separador in (",", ";", "\t"):
        campos = Counter(len(fila) for fila in csv.reader(lineas, delimiter=separador))
        if not campos:
            continue
        n_campos, n_lineas = max(campos.items(), key=lambda par: (par[1], par[0]))
        if n_campos >= 2 and (n_lineas, n_campos) > mejor_puntuacion:
            mejor, mejor_puntuacion = separador, (n_lineas, n_campos)
    return mejor


def _leer_csv(contenido: bytes, config: Configuracion) -> _Hoja:
    texto = _decodificar_csv(contenido)
    lector = csv.reader(io.StringIO(texto, newline=""), delimiter=_detectar_separador(texto))
    return _acumular(lector, config)


# --- Punto de entrada -----------------------------------------------------------------------------


_LECTORES = {"xlsx": _leer_xlsx, "xls": _leer_xls, "csv": _leer_csv}


def leer_archivo(nombre: str | None, contenido: bytes, config: Configuracion) -> TablaLeida:
    """Valida y lee el archivo completo desde memoria. Solo lanza `ErrorDataTab`."""
    nombre = limpiar_nombre(nombre)
    formato = _formato_por_extension(nombre)
    if not contenido:
        raise _vacio("El archivo está vacío.")
    try:
        hoja = _LECTORES[formato](contenido, config)
        return _construir(nombre, formato, hoja)
    except ErrorDataTab:
        raise
    except Exception:
        # El texto de la excepción podría incluir datos del archivo: no se propaga
        raise _ilegible() from None
