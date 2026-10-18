"""Genera los 3 archivos de prueba de esta carpeta. Todo es inventado y determinista (semilla fija).

Uso (desde la raíz, con las dependencias de backend/requirements-dev.txt):

    python datos-de-prueba/generar_datos_de_prueba.py
"""

import csv
import datetime as dt
import random
from pathlib import Path

import openpyxl
import xlwt

CARPETA = Path(__file__).resolve().parent
SEMILLA = 7


def elegir(azar, opciones):
    """opciones: lista de (valor, peso)."""
    return azar.choices([o for o, _ in opciones], weights=[p for _, p in opciones])[0]


def pacientes(azar, total=200):
    diagnosticos = [("Colecistitis aguda", 18), ("Colecistitis crónica", 14), ("Apendicitis aguda", 16), ("Hernia inguinal", 10),
                    ("Hernia umbilical", 6), ("Pancreatitis biliar", 5), ("Eventración", 3), ("Colelitiasis", 9),
                    ("Diverticulitis", 6), ("Reflujo gastroesofágico", 7), ("Obstrucción intestinal", 3), ("Quiste hepático", 3)]
    comorbilidades = [("Ninguna", 22), ("HTA", 14), ("Diabetes tipo 2", 10), ("Obesidad", 9), ("Asma", 5), ("Hipotiroidismo", 6),
                      ("HTA y diabetes", 8), ("Cardiopatía", 4), ("EPOC", 3), ("Insuficiencia renal", 3), ("Anemia", 5), ("Tabaquismo", 6)]
    ocupaciones = ["Docente", "Comerciante", "Agricultor", "Ama de casa", "Estudiante", "Conductor", "Obrero", "Jubilado",
                   "Enfermero", "Contador", "Cocinero", "Desempleado"]
    ciudades = [("Quito", 30), ("Guayaquil", 25), ("Cuenca", 12), ("Ambato", 8), ("Loja", 7), ("Manta", 6), ("Riobamba", 5), ("Ibarra", 4), ("Esmeraldas", 3)]
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Pacientes"
    hoja.append(["No.", "Sexo", "Edad", "Ocupación", "Ciudad", "Nivel educativo", "Diagnóstico", "Comorbilidad", "Cirugía",
                 "Peso (kg)", "Talla (cm)", "Dolor (0-10)", "Estancia (días)", "Fecha de ingreso", "Seguro", "Satisfacción"])
    inicio = dt.date(2024, 1, 8)
    for n in range(1, total + 1):
        mujer = azar.random() < 0.58
        edad = int(min(92, max(16, azar.gauss(47, 17))))
        talla = round(azar.gauss(160 if mujer else 171, 7))
        peso = round(azar.gauss(68 if mujer else 78, 13), 1)
        cirugia = elegir(azar, [("Laparoscópica", 62), ("Abierta", 24), ("Ambulatoria", 14)])
        hoja.append([
            n, "Femenino" if mujer else "Masculino", edad if azar.random() > 0.03 else None,
            azar.choice(ocupaciones), elegir(azar, ciudades),
            elegir(azar, [("Primaria", 22), ("Secundaria", 40), ("Superior", 30), ("Posgrado", 8)]),
            elegir(azar, diagnosticos), elegir(azar, comorbilidades), cirugia,
            peso, talla, azar.randint(0, 10), max(1, int(azar.gauss(3 if cirugia != "Abierta" else 6, 2))),
            inicio + dt.timedelta(days=azar.randint(0, 280)),
            elegir(azar, [("Público", 52), ("Privado", 30), ("Mixto", 12), ("Sin seguro", 6)]),
            elegir(azar, [("Muy satisfecho", 38), ("Satisfecho", 36), ("Neutral", 14), ("Insatisfecho", 9), ("Muy insatisfecho", 3)]),
        ])
    for fila in hoja.iter_rows(min_row=2, min_col=14, max_col=14):
        fila[0].number_format = "DD/MM/YYYY"
    return libro


def encuesta(azar, ruta, total=120):
    """CSV de Excel en español: separador ;, codificación Windows-1252. Datos de una encuesta escolar, no médicos."""
    materias = ["Matemáticas", "Lengua", "Ciencias", "Historia", "Inglés", "Educación física", "Arte"]
    with open(ruta, "w", newline="", encoding="cp1252") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Estudiante", "Curso", "Edad", "Materia favorita", "Horas de estudio por semana", "Transporte", "Nota final", "Usa celular en clase"])
        for n in range(1, total + 1):
            w.writerow([n, azar.choice(["1ro", "2do", "3ro", "4to", "5to"]), azar.randint(12, 18), azar.choice(materias),
                        azar.randint(0, 20), elegir(azar, [("Bus", 40), ("A pie", 25), ("Bicicleta", 8), ("Auto familiar", 22), ("Moto", 5)]),
                        round(min(10, max(3, azar.gauss(7.4, 1.5))), 1), azar.choice(["Sí", "No", "A veces"])])


def ventas(azar, ruta, total=90):
    """Formato antiguo .xls con ventas de una tienda."""
    productos = [("Cuaderno", 2.5), ("Lápiz", 0.6), ("Mochila", 24.0), ("Calculadora", 15.0), ("Marcadores", 5.5), ("Regla", 1.2), ("Carpeta", 3.2)]
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("Ventas")
    for c, t in enumerate(["Factura", "Mes", "Producto", "Cantidad", "Precio unitario", "Sucursal", "Forma de pago", "Cliente nuevo"]):
        hoja.write(0, c, t)
    meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio"]
    for n in range(1, total + 1):
        producto, precio = azar.choice(productos)
        fila = [1000 + n, azar.choice(meses), producto, azar.randint(1, 12), precio,
                azar.choice(["Centro", "Norte", "Sur"]), elegir(azar, [("Efectivo", 45), ("Tarjeta", 40), ("Transferencia", 15)]),
                azar.choice(["Sí", "No"])]
        for c, v in enumerate(fila):
            hoja.write(n, c, v)
    libro.save(str(ruta))


def main():
    destino = CARPETA / "pacientes-ficticios.xlsx"
    libro = pacientes(random.Random(SEMILLA))
    libro.properties.created = libro.properties.modified = dt.datetime(2024, 10, 18, 12, 0, 0)
    libro.save(destino)
    encuesta(random.Random(SEMILLA + 1), CARPETA / "encuesta-escolar.csv")
    ventas(random.Random(SEMILLA + 2), CARPETA / "ventas-tienda.xls")


if __name__ == "__main__":
    main()
